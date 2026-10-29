"""Pure SPL parser and HTTP-contract tests; no live SPL or site required."""

import importlib.util
import io
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1] / "factory_management" / "spl_po"


def load(name, filename):
    spec = importlib.util.spec_from_file_location(name, ROOT / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


parser = load("spl_parser", "parser.py")


def fixture_xlsx(rows):
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "ExcelOut(PO)"
    sheet.append(["SPL report"])
    sheet.append(list(parser.REQUIRED) + ["Revision No.", "Year", "Planning Ssn", "Approval Date"])
    for row in rows:
        sheet.append(row)
    data = io.BytesIO()
    book.save(data)
    return data.getvalue()


class ParserTests(unittest.TestCase):
    row = ["PO-1", "I-1", "Item", "07ABC", "09", "Red", "M", "Medium", 2.5,
           "R1", 2026, "SS", "2026/10/29"]

    def test_xlsx_normalizes_rows_dates_and_rounding(self):
        rows = parser.parse_workbook(fixture_xlsx([self.row]), "report.xlsx")
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["market_code"], "07")
        self.assertEqual(rows[0]["sample_code"], "ABC")
        self.assertEqual(rows[0]["order_qty"], 3)
        self.assertEqual(rows[0]["approval_date"], "2026-10-29")

    def test_xls_reader(self):
        import xlwt
        book = xlwt.Workbook()
        sheet = book.add_sheet("ExcelOut(PO)")
        for column, value in enumerate(parser.REQUIRED):
            sheet.write(1, column, value)
        for column, value in enumerate(self.row[:len(parser.REQUIRED)]):
            sheet.write(2, column, value)
        data = io.BytesIO()
        book.save(data)
        rows = parser.parse_workbook(data.getvalue(), "report.xls")
        self.assertEqual(rows[0]["order_qty"], 3)

    def test_missing_column_fails_with_code(self):
        content = fixture_xlsx([self.row])
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(content))
        book.active.cell(2, 1).value = "Unknown"
        data = io.BytesIO()
        book.save(data)
        with self.assertRaises(parser.WorkbookError) as caught:
            parser.parse_workbook(data.getvalue(), "report.xlsx")
        self.assertEqual(caught.exception.code, "SPL_COLUMNS_MISSING")

    def test_invalid_date_and_missing_sheet_are_rejected(self):
        row = self.row.copy()
        row[-1] = "32/99/2026"
        with self.assertRaises(parser.WorkbookError) as caught:
            parser.parse_workbook(fixture_xlsx([row]), "report.xlsx")
        self.assertEqual(caught.exception.code, "SPL_INVALID_DATE")
        import openpyxl
        book = openpyxl.Workbook()
        data = io.BytesIO()
        book.save(data)
        with self.assertRaises(parser.WorkbookError) as caught:
            parser.parse_workbook(data.getvalue(), "report.xlsx")
        self.assertEqual(caught.exception.code, "SPL_SHEET_MISSING")

    def test_invalid_sample_and_size_guard(self):
        bad = self.row.copy()
        bad[3] = "A"
        with self.assertRaises(parser.WorkbookError) as caught:
            parser.parse_workbook(fixture_xlsx([bad]), "report.xlsx")
        self.assertEqual(caught.exception.code, "SPL_SAMPLE_INVALID")
        with self.assertRaises(parser.WorkbookError) as caught:
            parser.parse_workbook(b"x" * (parser.MAX_BYTES + 1), "report.xlsx")
        self.assertEqual(caught.exception.code, "SPL_FILE_SIZE_INVALID")

    def test_price_and_currency_follow_legacy_importer(self):
        self.assertEqual(str(parser.parse_unit_price("USD 1.234,56")), "1234.5600")
        self.assertEqual(str(parser.parse_unit_price("1,234.56")), "1234.5600")
        self.assertEqual(parser.normalize_currency("Japanese Yen"), "JPY")


class GatewayContractTests(unittest.TestCase):
    def test_http_auth_failure_is_redacted(self):
        frappe_stub = types.ModuleType("frappe")
        frappe_stub.conf = {"spl_po_allowed_hosts": ["spl.fastretailing.com"]}
        sys.modules["frappe"] = frappe_stub
        package = types.ModuleType("factory_management")
        package.__path__ = []
        subpackage = types.ModuleType("factory_management.spl_po")
        subpackage.__path__ = []
        sys.modules["factory_management"] = package
        sys.modules["factory_management.spl_po"] = subpackage
        sys.modules["factory_management.spl_po.parser"] = parser
        gateway = load("spl_gateway_auth", "gateway.py")
        class FakeSession:
            def request(self, *args, **kwargs):
                return types.SimpleNamespace(status_code=401, text="secret-token")
        with self.assertRaises(gateway.GatewayError) as caught:
            gateway._request(FakeSession(), "POST", "https://spl.fastretailing.com/api/mdm/auth/login", 10)
        self.assertEqual(str(caught.exception), "SPL_HTTP_FAILED")

    def test_insecure_host_is_rejected(self):
        frappe_stub = types.ModuleType("frappe")
        frappe_stub.conf = {"spl_po_allowed_hosts": ["spl.fastretailing.com"]}
        sys.modules["frappe"] = frappe_stub
        package = types.ModuleType("factory_management")
        package.__path__ = []
        subpackage = types.ModuleType("factory_management.spl_po")
        subpackage.__path__ = []
        sys.modules["factory_management"] = package
        sys.modules["factory_management.spl_po"] = subpackage
        sys.modules["factory_management.spl_po.parser"] = parser
        gateway = load("spl_gateway_reject", "gateway.py")
        with self.assertRaises(gateway.GatewayError) as caught:
            gateway._origin(types.SimpleNamespace(base_url="http://spl.fastretailing.com"))
        self.assertEqual(str(caught.exception), "SPL_HOST_INVALID")

    def test_login_report_and_download_use_certificate_and_multipart(self):
        # Stub only Frappe's configuration; no site or credentials are involved.
        frappe_stub = types.ModuleType("frappe")
        frappe_stub.conf = {"spl_po_allowed_hosts": ["spl.fastretailing.com"]}
        package = types.ModuleType("factory_management")
        package.__path__ = []
        subpackage = types.ModuleType("factory_management.spl_po")
        subpackage.__path__ = []
        sys.modules["frappe"] = frappe_stub
        sys.modules["factory_management"] = package
        sys.modules["factory_management.spl_po"] = subpackage
        sys.modules["factory_management.spl_po.parser"] = parser
        gateway = load("spl_gateway", "gateway.py")
        class Settings:
            base_url = "https://spl.fastretailing.com"
            spl_user_id = "test"
            request_timeout_seconds = 120
            mode = "01"
            answer_status_id = "0003"
            item_brand_id = "010"
            year_season_ids = "20261"
            management_factory_id = "015J"
            document_status_ids = "0001"
            report_offset = 0
            def get_password(self, field):
                return "placeholder"
        class Response:
            status_code = 200
            def __init__(self, payload=None):
                self.payload = payload
            def json(self):
                return self.payload
            def iter_content(self, size):
                yield b"data"
        class Session:
            cert = None
            calls = []
            def __enter__(self): return self
            def __exit__(self, *args): return None
            def request(self, method, url, **kwargs):
                self.calls.append((method, url, kwargs, self.cert))
                if url.endswith("login"): return Response({"accessToken": "fake-token"})
                if url.endswith("reportxlsxbaseoncondition"):
                    return Response({"filePath": "reports/po.xlsx"})
                return Response()
        from contextlib import contextmanager
        @contextmanager
        def fake_pair(settings):
            yield ("cert.pem", "key.pem")
        session = Session()
        with patch.object(gateway, "_pem_pair", fake_pair), patch.object(gateway.requests, "Session", return_value=session):
            content, path = gateway.fetch_report(Settings())
        self.assertEqual((content, path), (b"data", "reports/po.xlsx"))
        self.assertEqual([call[0] for call in session.calls], ["POST", "POST", "GET"])
        self.assertEqual(session.calls[1][2]["files"][0], ("mode", (None, "01")))
        self.assertEqual(session.calls[1][3], ("cert.pem", "key.pem"))
        self.assertEqual(session.calls[2][2]["cookies"], {"accessToken": "fake-token"})


if __name__ == "__main__":
    unittest.main()
