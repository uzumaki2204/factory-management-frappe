"""Streaming SPL ExcelOut(PO) parser; no Frappe dependency."""

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_EVEN, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

SHEET = "ExcelOut(PO)"
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 50_000
REQUIRED = (
    "Order No.", "Item Code", "Item", "Sample Code", "Color Code",
    "Color", "Size Code", "Size", "Order Qty(pcs)",
)
TEXT = {
    "po_number": "Order No.", "revision_no": "Revision No.",
    "planning_season_code": "Planning Ssn", "global_business_unit": "Global Business Unit",
    "business_unit": "Business Unit", "item_brand": "Item Brand",
    "department": "Department", "document_status_external": "Document Status",
    "approver": "Approver", "answered_status_external": "Answered Status",
    "sales_contract_no": "Sales Contract No", "bill_to_name": "Bill to",
    "bill_to_address": "Address", "bill_to_telephone": "Telephone",
    "bill_to_person_in_charge": "Person in Charge", "incoterms": "Incoterms",
    "transportation_method": "Transportation Method", "payment_terms": "Payment Terms",
    "payment_currency": "Payment Currency", "order_conditions": "Order Conditions",
    "order_plan_number": "Order Plan Number", "item_code": "Item Code",
    "item_name": "Item", "representative_sample_code": "Representative Sample Code",
    "color_code": "Color Code", "color_name": "Color", "size_code": "Size Code",
    "size_name": "Size", "pattern_dimension_code": "Pattern Dimension Code",
    "sku": "SKU Code", "sku_code_for_shipping": "SKU Code (For Shipping)",
    "sample_code_raw": "Sample Code", "external_remark": "Remark",
    "allocation_order_number": "Allocation Order Number",
    "original_sales_price_text": "Original Sales Price",
}
DATES = {
    "revised_date": "Revised Date", "approval_date": "Approval Date",
    "deadline_date": "Deadline", "decision_date": "Decision Date",
    "contracted_etd": "Contracted ETD", "eta_wh": "ETA WH",
}


class WorkbookError(ValueError):
    """An invalid SPL report that is safe to show to the operator."""

    def __init__(self, code, **context):
        self.code = code
        self.context = context
        super().__init__(code)


def parse_unit_price(raw):
    value = _text(raw)
    if not value:
        return Decimal("0")
    try:
        parsed = Decimal(value.replace(",", ""))
    except InvalidOperation:
        filtered = "".join(char for char in value if char.isdigit() or char in ".,")
        if not filtered:
            return Decimal("0")
        last_dot, last_comma = filtered.rfind("."), filtered.rfind(",")
        if last_dot >= 0 and last_comma >= 0:
            normalized = (filtered.replace(",", "") if last_dot > last_comma
                          else filtered.replace(".", "").replace(",", "."))
        else:
            normalized = filtered.replace(",", ".")
        try:
            parsed = Decimal(normalized)
        except InvalidOperation:
            return Decimal("0")
    try:
        return parsed.quantize(Decimal("0.0001"), rounding=ROUND_HALF_EVEN)
    except InvalidOperation:
        return Decimal("0")


def normalize_currency(raw):
    value = _text(raw)
    return {"US Dollar": "USD", "Japanese Yen": "JPY", "Vietnamese Dong": "VND"}.get(value, value[:10])


def _text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _date(value, row, header):
    if value is None or _text(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (float, int)):
        try:
            return (date(1899, 12, 30) + timedelta(days=float(value))).isoformat()
        except (OverflowError, ValueError):
            pass
    raw = _text(value)
    for pattern in ("%Y/%m/%d", "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    raise WorkbookError("SPL_INVALID_DATE", row=row, column=header)


def _integer(value, row, header, required=False):
    raw = _text(value).replace(",", "")
    if not raw:
        if required:
            raise WorkbookError("SPL_VALUE_REQUIRED", row=row, column=header)
        return None
    try:
        return int(Decimal(raw).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        raise WorkbookError("SPL_INVALID_NUMBER", row=row, column=header) from None


def _rows_xlsx(stream):
    from openpyxl import load_workbook

    book = load_workbook(stream, read_only=True, data_only=True)
    try:
        sheet = next((book[name] for name in book.sheetnames if name.lower() == SHEET.lower()), None)
        if sheet is None:
            raise WorkbookError("SPL_SHEET_MISSING", sheet=SHEET)
        yield from sheet.iter_rows(values_only=True)
    finally:
        book.close()


def _rows_xls(stream):
    import xlrd

    book = xlrd.open_workbook(file_contents=stream.read(), on_demand=True)
    try:
        sheet = next((book.sheet_by_name(name) for name in book.sheet_names() if name.lower() == SHEET.lower()), None)
        if sheet is None:
            raise WorkbookError("SPL_SHEET_MISSING", sheet=SHEET)
        for index in range(sheet.nrows):
            yield sheet.row_values(index)
    finally:
        book.release_resources()


def parse_workbook(content, filename):
    if not content or len(content) > MAX_BYTES:
        raise WorkbookError("SPL_FILE_SIZE_INVALID")
    suffix = Path(filename).suffix.lower()
    if suffix not in (".xls", ".xlsx"):
        raise WorkbookError("SPL_FILE_FORMAT_INVALID")
    source = _rows_xlsx(BytesIO(content)) if suffix == ".xlsx" else _rows_xls(BytesIO(content))
    iterator = iter(source)
    try:
        next(iterator)  # SPL report caption row.
        headers = {}
        for index, value in enumerate(next(iterator, ())):
            header = " ".join(_text(value).split())
            if header:
                headers.setdefault(header, index)
        missing = [name for name in REQUIRED if name not in headers]
        if missing:
            raise WorkbookError("SPL_COLUMNS_MISSING", columns=", ".join(missing))
        rows = []
        for row_number, values in enumerate(iterator, start=3):
            def cell(name):
                index = headers.get(name)
                return values[index] if index is not None and index < len(values) else None
            if not _text(cell("Order No.")):
                continue
            if len(rows) >= MAX_ROWS:
                raise WorkbookError("SPL_ROW_LIMIT")
            record = {name: _text(cell(header)) for name, header in TEXT.items()}
            for required in REQUIRED:
                if not _text(cell(required)):
                    raise WorkbookError("SPL_VALUE_REQUIRED", row=row_number, column=required)
            record.update({name: _date(cell(header), row_number, header)
                           for name, header in DATES.items()})
            record["planning_year"] = _integer(cell("Year"), row_number, "Year")
            record["order_qty"] = _integer(cell("Order Qty(pcs)"), row_number, "Order Qty(pcs)", True)
            record["row_number"] = row_number
            sample = record["sample_code_raw"].strip()
            if len(sample) < 3:
                raise WorkbookError("SPL_SAMPLE_INVALID", row=row_number)
            record["market_code"] = sample[:2]
            record["sample_code"] = sample[2:]
            rows.append(record)
        return rows
    except WorkbookError:
        raise
    except Exception:
        raise WorkbookError("SPL_FILE_FORMAT_INVALID") from None
    finally:
        close = getattr(iterator, "close", None)
        if close:
            close()
